from django.contrib import messages
from django.contrib.auth.mixins import LoginRequiredMixin, PermissionRequiredMixin
from django.urls import reverse_lazy
from django.contrib.auth.views import redirect_to_login
from django.views.generic import (
    CreateView,
    DeleteView,
    DetailView,
    ListView,
    UpdateView,
)
from .models import Producto
from .forms import ProductoForm
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from django.http import HttpResponse
from django.contrib.auth.decorators import login_required


class ProductoListView(ListView):
    """Lista paginada de productos activos."""

    model = Producto
    template_name = 'productos/lista.html'
    context_object_name = 'productos'
    paginate_by = 20

    def get_queryset(self):
        """Devuelve productos activos con FKs pre-cargadas."""
        return (
            Producto.objects
            .select_related('categoria', 'proveedor')
            .filter(activo=True)
            .order_by('nombre')
        )


class ProductoDetailView(DetailView):
    """Detalle de un producto por ID de URL."""

    model = Producto
    template_name = 'productos/detalle.html'
    context_object_name = 'producto'
    pk_url_kwarg = 'producto_id'

    def get_queryset(self):
        """Pre-carga FKs para evitar consultas adicionales."""
        return Producto.objects.select_related(
            'categoria',
            'proveedor'
        )


class ProductoCreateView(LoginRequiredMixin, CreateView):
    """Formulario para crear un producto nuevo."""

    model = Producto
    form_class = ProductoForm
    template_name = 'productos/crear.html'
    success_url = reverse_lazy('productos:lista')

    def form_valid(self, form):
        """Guarda el producto y muestra un mensaje de éxito."""
        response = super().form_valid(form)

        messages.success(
            self.request,
            f'Producto "{self.object.nombre}" creado exitosamente.'
        )

        return response


class ProductoUpdateView(LoginRequiredMixin, UpdateView):
    """Formulario para editar un producto existente."""

    model = Producto
    form_class = ProductoForm
    template_name = 'productos/editar.html'
    pk_url_kwarg = 'producto_id'
    success_url = reverse_lazy('productos:lista')

    def form_valid(self, form):
        """Guarda los cambios y muestra un mensaje de éxito."""
        response = super().form_valid(form)

        messages.success(
            self.request,
            f'Producto "{self.object.nombre}" actualizado.'
        )

        return response


class ProductoDeleteView(
    LoginRequiredMixin,
    PermissionRequiredMixin,
    DeleteView
):
    """Confirmación y borrado de un producto.

    Requiere el permiso:
    productos.delete_producto
    """

    model = Producto
    template_name = 'productos/eliminar.html'
    pk_url_kwarg = 'producto_id'
    success_url = reverse_lazy('productos:lista')
    permission_required = 'productos.delete_producto'
    raise_exception = True

    def handle_no_permission(self):
        """Redirige al login si no está autenticado; si no, responde 403."""
        if not self.request.user.is_authenticated:
            return redirect_to_login(
                self.request.get_full_path(),
                self.get_login_url(),
                self.get_redirect_field_name()
            )

        return super().handle_no_permission()

    def form_valid(self, form):
        """Elimina el producto y muestra un mensaje de éxito."""
        nombre = self.object.nombre

        response = super().form_valid(form)

        messages.success(
            self.request,
            f'Producto "{nombre}" eliminado.'
        )

        return response
@login_required
def exportar_excel(request) -> HttpResponse:
    """Genera y descarga un archivo Excel (.xlsx) con el catálogo de productos."""

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Productos"

    headers = ['ID', 'Nombre', 'Categoría', 'Precio', 'Stock', 'Activo']
    ws.append(headers)

    fill_navy = PatternFill(
        start_color='0A2342',
        end_color='0A2342',
        fill_type='solid'
    )

    font_header = Font(
        name='Arial',
        size=11,
        bold=True,
        color='FFFFFF'
    )

    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = fill_navy
        cell.font = font_header
        cell.alignment = Alignment(
            horizontal='center',
            vertical='center'
        )

    productos = Producto.objects.select_related('categoria').all()

    for p in productos:
        ws.append([
            p.pk,
            p.nombre,
            p.categoria.nombre if p.categoria else 'Sin Categoría',
            float(p.precio),
            p.stock,
            'Sí' if p.activo else 'No',
        ])

    for col in ws.columns:
        max_len = max(
            len(str(cell.value or ''))
            for cell in col
        )

        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

    response = HttpResponse(
        content_type=(
            'application/vnd.openxmlformats-officedocument.'
            'spreadsheetml.sheet'
        )
    )

    response['Content-Disposition'] = (
        'attachment; filename="catalogo_productos.xlsx"'
    )

    wb.save(response)

    return response


